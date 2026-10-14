import os
import re
import json
import urllib.request
import subprocess
import smtplib
from email.message import EmailMessage
from datetime import datetime, timezone, timedelta
import streamlit as st
import docx
import openpyxl
from kannada_converter import ascii_to_unicode

# Page configuration
st.set_page_config(
    page_title="Kannada Unicode Converter", page_icon="🔤", layout="centered"
)

st.title("Kannada ASCII to Unicode Converter")
st.markdown("**Developed by Nethen**")
st.write(
    "Upload your `.doc`, `.docx`, or `.xlsx` file to convert it to Unicode format."
)


# --- SILENT VISITOR TRACKER & EMAIL ALERT ---
def get_visitor_info():
    """Captures visitor IP, approximate location/ISP, device, and timestamp silently."""
    headers = getattr(st.context, "headers", {}) or {}

    # 1. Get real client IP (prioritize X-Forwarded-For on Streamlit Cloud)
    forwarded = headers.get("X-Forwarded-For", "")
    if forwarded:
        ip = forwarded.split(",")[0].strip()
    else:
        ip = getattr(st.context, "ip_address", None) or "Unknown"

    # 2. Browser & System details
    user_agent = headers.get("User-Agent", "Unknown Device/Browser")
    platform = headers.get("Sec-Ch-Ua-Platform", "Unknown OS").strip('"')
    tz = getattr(st.context, "timezone", None) or "Unknown Timezone"
    locale = getattr(st.context, "locale", None) or headers.get(
        "Accept-Language", "Unknown"
    )

    # 3. Geolocation lookup from IP
    location_str = "Location unavailable"
    isp_str = "Unknown ISP"
    if ip and ip not in ["Unknown", "127.0.0.1", "::1"]:
        try:
            url = f"http://ip-api.com/json/{ip}?fields=status,country,regionName,city,isp"
            req = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0"})
            with urllib.request.urlopen(req, timeout=2) as resp:
                data = json.loads(resp.read().decode("utf-8"))
                if data.get("status") == "success":
                    location_str = f"{data.get('city', '')}, {data.get('regionName', '')}, {data.get('country', '')}"
                    isp_str = data.get("isp", "Unknown ISP")
        except Exception:
            pass

    ist_now = datetime.now(timezone(timedelta(hours=5, minutes=30))).strftime(
        "%Y-%m-%d %I:%M:%S %p IST"
    )

    return (
        f"• Time: {ist_now}\n"
        f"• IP Address: {ip}\n"
        f"• Location: {location_str}\n"
        f"• ISP / Network: {isp_str}\n"
        f"• Platform: {platform}\n"
        f"• Timezone: {tz}\n"
        f"• Language: {locale}\n"
        f"• User Agent: {user_agent}\n"
    )


def notify_visit(event_label="Opened the App"):
    """Prints visitor details to Streamlit console logs and sends an email alert."""
    try:
        info = get_visitor_info()
        print(
            f"\n--- 🔔 VISITOR TRACKER ({event_label}) ---\n{info}------------------------------------------",
            flush=True,
        )

        sender_email = st.secrets.get("NOTIFY_EMAIL")
        app_password = st.secrets.get("NOTIFY_APP_PASSWORD")
        receiver_email = st.secrets.get("RECEIVER_EMAIL", "dsdanethen@gmail.com")

        if sender_email and app_password:
            msg = EmailMessage()
            msg["Subject"] = f"🔔 Kannada Converter: {event_label}"
            msg["From"] = sender_email
            msg["To"] = receiver_email
            msg.set_content(
                f"Hello Nethen,\n\nActivity detected on your Kannada ASCII to Unicode Converter ({event_label}):\n\n{info}"
            )

            with smtplib.SMTP_SSL("smtp.gmail.com", 465) as server:
                server.login(sender_email, app_password)
                server.send_message(msg)
    except Exception:
        pass


# Trigger once per session as soon as someone opens the link
if "visit_tracked" not in st.session_state:
    notify_visit("Opened the App Link")
    st.session_state.visit_tracked = True
# --- END TRACKER ---


def transliterate_text_safely(text):
    if not isinstance(text, str) or not text.strip():
        return text

    email_pattern = r"[a-zA-Z0-9_.+-]+@[a-zA-Z0-9-]+\.[a-zA-Z0-9-.]+"
    url_pattern = r"https?://\S+|www\.\S+"

    emails = re.findall(email_pattern, text)
    urls = re.findall(url_pattern, text)

    protected_items = emails + urls
    temp_text = text

    for i, item in enumerate(protected_items):
        temp_text = temp_text.replace(item, f"_PROTECTED{i}__")

    tokens = re.split(r"(\s+)", temp_text)

    translated_tokens = []
    for token in tokens:
        if token.isspace() or token.startswith("_PROTECTED"):
            translated_tokens.append(token)
        else:
            try:
                conv = ascii_to_unicode(token)
                translated_tokens.append(conv if conv else token)
            except Exception:
                translated_tokens.append(token)

    converted = "".join(translated_tokens)

    for i, item in enumerate(protected_items):
        converted = converted.replace(f"_PROTECTED{i}__", item)

    manual_fixes = {"¯ï": "ಲ್", "¯i": "ಲ್"}

    for ascii_bug, unicode_fix in manual_fixes.items():
        converted = converted.replace(ascii_bug, unicode_fix)

    return converted


def convert_doc_to_docx(input_file):
    subprocess.run(
        ["soffice", "--headless", "--convert-to", "docx", input_file],
        stdout=subprocess.DEVNULL,
        stderr=subprocess.DEVNULL,
    )
    output_file = f"{os.path.splitext(input_file)[0]}.docx"
    return output_file if os.path.exists(output_file) else None


def process_word_document(input_path, output_path):
    doc = docx.Document(input_path)

    for para in doc.paragraphs:
        if para.text.strip():
            translated = transliterate_text_safely(para.text)
            para.clear()
            new_run = para.add_run(translated)
            new_run.font.name = "Tunga"

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    if para.text.strip():
                        translated = transliterate_text_safely(para.text)
                        para.clear()
                        new_run = para.add_run(translated)
                        new_run.font.name = "Tunga"

    doc.save(output_path)


def process_excel_document(input_path, output_path):
    wb = openpyxl.load_workbook(input_path)

    for sheet in wb.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.strip():
                    cell.value = transliterate_text_safely(cell.value)

    wb.save(output_path)


# Streamlit File Uploader Widget
uploaded_file = st.file_uploader(
    "ಫೈಲ್ ಅನ್ನು ಆಯ್ಕೆ ಮಾಡಿ (Choose file)", type=["doc", "docx", "xlsx"]
)

if uploaded_file is not None:
    filename = uploaded_file.name

    # Log and email when a file is uploaded
    if st.session_state.get("last_tracked_file") != filename:
        notify_visit(f"Uploaded File: {filename}")
        st.session_state.last_tracked_file = filename

    # Save uploaded file locally in the container
    with open(filename, "wb") as f:
        f.write(uploaded_file.getbuffer())

    file_ext = os.path.splitext(filename)[1].lower()
    final_output_name = f"kannada_unicode_{os.path.splitext(filename)[0]}"
    final_output = None

    with st.spinner("⚙️ Processing and translating text... Please wait."):
        if file_ext == ".doc":
            temp_docx = convert_doc_to_docx(filename)
            if temp_docx:
                final_output = f"{final_output_name}.docx"
                process_word_document(temp_docx, final_output)
        elif file_ext == ".docx":
            final_output = f"{final_output_name}.docx"
            process_word_document(filename, final_output)
        elif file_ext == ".xlsx":
            final_output = f"{final_output_name}.xlsx"
            process_excel_document(filename, final_output)

    if final_output and os.path.exists(final_output):
        st.success("✅ Processing complete!")
        with open(final_output, "rb") as file:
            st.download_button(
                label="📥 Download Converted File",
                data=file,
                file_name=final_output,
                mime=(
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
                    if file_ext in [".doc", ".docx"]
                    else "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
                ),
            )

        # Cleanup temporary files
        try:
            os.remove(filename)
            if file_ext == ".doc" and os.path.exists(temp_docx):
                os.remove(temp_docx)
        except Exception:
            pass
